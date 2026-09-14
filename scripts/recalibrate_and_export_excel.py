#!/usr/bin/env python3
"""
task-12 (افزونه) — بازکالیبره آستانه‌های توصیه + تولید Excel کامل.

مشکل نسخه قبلی: آستانه‌های ثابت 0.4/0.7 کاملاً بالاتر از دامنه واقعی نمره‌ها
[-0.03, 0.35] بودند → هر ۵۷۴ کاندید «investigate» شده بودند.
راه‌حل: آستانه‌های نسبی بر پایه صدک (select ≥ صدک ۹۰، monitor ≥ صدک ۶۰).

خروجی‌ها:
  - 05-Data/processed/export-candidates-ranked-recalibrated.parquet
  - 07-Exports/iran-export-candidates-500.xlsx (۷ شیت)
  - 06-Analysis/export-candidates/_executive-ranking.md (به‌روزرسانی با Top 50)

نکته پیاده‌سازی شیت By-Target-Country: به‌جای نسبت‌دادن total_value_1404 کلِ HS به
هر کشور (که با ۵ کشورِ top-5 تا ۵ برابر تورش می‌دهد)، ارزش واقعی هر جفت
HS×Country از trend-analysis-5y.parquet (task-10) خوانده می‌شود — پوشش join: ۱۰۰٪.
"""
from __future__ import annotations

from datetime import date
from pathlib import Path

import numpy as np
import pandas as pd
import pyarrow as pa
import pyarrow.parquet as pq
from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

REPO_ROOT = Path(__file__).resolve().parents[1]
INPUT_PATH = REPO_ROOT / "05-Data/processed/export-candidates-ranked.parquet"
TREND_PATH = REPO_ROOT / "05-Data/processed/trend-analysis-5y.parquet"
OUTPUT_PARQUET = REPO_ROOT / "05-Data/processed/export-candidates-ranked-recalibrated.parquet"
OUTPUT_EXCEL = REPO_ROOT / "07-Exports/iran-export-candidates-500.xlsx"
OUTPUT_EXEC = REPO_ROOT / "06-Analysis/export-candidates/_executive-ranking.md"

# آستانه‌های صدک
SELECT_PERCENTILE = 90
MONITOR_PERCENTILE = 60

TODAY = date.today().isoformat()
TODAY_FA = "1405-06-23"

# ─── استایل‌ها (طبق اسکریپت مرجع کاربر) ───
HEADER_FONT = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
BORDER = Border(
    left=Side(style="thin", color="CCCCCC"),
    right=Side(style="thin", color="CCCCCC"),
    top=Side(style="thin", color="CCCCCC"),
    bottom=Side(style="thin", color="CCCCCC"),
)

# رنگ‌های وضعیت توصیه (پالت design-system: accent سبز/کهربایی با پس‌زمینه روشن)
CF_SELECT_FILL = PatternFill(start_color="E8F5E9", end_color="E8F5E9", fill_type="solid")
CF_SELECT_FONT = Font(name="Calibri", color="1B7D46", bold=True)
CF_MONITOR_FILL = PatternFill(start_color="FEF9E7", end_color="FEF9E7", fill_type="solid")
CF_MONITOR_FONT = Font(name="Calibri", color="D4820A", bold=True)

# فرمت اعداد به ازای نام ستون
NUMBER_FORMATS: dict[str, str] = {
    "rank": "0",
    "export_score": "0.000",
    "cagr_5y_aggregated": "0.0%",
    "slope_aggregated": "#,##0",
    "mean_recent_3y_aggregated": "#,##0",
    "n_destinations_active": "0",
    "trend_consistency_aggregated": "0.00",
    "cv_aggregated": "0.00",
    "total_value_5y": "#,##0",
    "total_value_1404": "#,##0",
    "n_strong_growth": "0",
    "n_moderate_growth": "0",
    "n_emerging": "0",
    "n_declining": "0",
    "n_disappearing": "0",
    "max_country_share": "0.0%",
    "score_growth": "0.000",
    "score_slope": "0.000",
    "score_volume": "0.000",
    "score_diversity": "0.000",
    "score_consistency": "0.000",
    "score_volatility_penalty": "0.000",
    # شیت‌های ۵ و ۶
    "n_candidates": "0",
    "n_select": "0",
    "n_monitor": "0",
    "total_score": "0.00",
    "avg_score": "0.000",
    "max_score": "0.000",
    "n_hs": "0",
}

# ستون‌های Excel (به ترتیب منطقی)
EXCEL_COLS = [
    "rank", "hs_code", "hs_description", "hs_chapter",
    "export_score", "recommendation",
    "cagr_5y_aggregated", "slope_aggregated", "mean_recent_3y_aggregated",
    "n_destinations_active", "trend_consistency_aggregated", "cv_aggregated",
    "total_value_5y", "total_value_1404",
    "n_strong_growth", "n_moderate_growth", "n_emerging",
    "n_declining", "n_disappearing",
    "top_5_countries", "max_country_share",
    "risk_flags",
    "score_growth", "score_slope", "score_volume",
    "score_diversity", "score_consistency", "score_volatility_penalty",
]


def _native_value(val):
    """numpy → نوع بومی پایتون؛ NaN/NaT → None (سلول خالی، نه رشته خالی)."""
    if isinstance(val, np.integer):
        return int(val)
    if isinstance(val, np.floating):
        return None if np.isnan(val) else float(val)
    if val is None:
        return None
    try:
        if pd.isna(val):
            return None
    except (TypeError, ValueError):
        pass
    return val


def _style_header(ws, n_cols: int) -> None:
    for col in range(1, n_cols + 1):
        cell = ws.cell(row=1, column=col)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = HEADER_ALIGN
        cell.border = BORDER
    ws.row_dimensions[1].height = 30


def _autofit_columns(ws, max_width: int = 50) -> None:
    for column in ws.columns:
        try:
            col_letter = get_column_letter(column[0].column)
        except (AttributeError, IndexError):
            continue
        max_length = 0
        for cell in column:
            if cell.value is not None:
                length = len(str(cell.value))
                if length > max_length:
                    max_length = length
        ws.column_dimensions[col_letter].width = min(max_length + 2, max_width)


def write_table(ws, frame: pd.DataFrame, cols: list[str]) -> int:
    """هدر (ردیف ۱) + داده + استایل + عرض ستون + AutoFilter + Freeze. تعداد ردیف داده را برمی‌گرداند."""
    for col_idx, col_name in enumerate(cols, 1):
        ws.cell(row=1, column=col_idx, value=col_name)
    fmt_map = NUMBER_FORMATS
    for row_idx, (_, row) in enumerate(frame.iterrows(), 2):
        for col_idx, col_name in enumerate(cols, 1):
            val = _native_value(row.get(col_name))
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            fmt = fmt_map.get(col_name)
            if fmt is not None and val is not None:
                cell.number_format = fmt
    _style_header(ws, len(cols))
    _autofit_columns(ws)
    ws.auto_filter.ref = ws.dimensions
    ws.freeze_panes = "A2"
    return len(frame)


def add_conditional_formatting(ws, cols: list[str], n_rows: int) -> None:
    """Color scale روی export_score + رنگ وضعیت recommendation (در صورت وجود ستون)."""
    if n_rows <= 0:
        return
    if "export_score" in cols:
        letter = get_column_letter(cols.index("export_score") + 1)
        rng = f"{letter}2:{letter}{n_rows + 1}"
        ws.conditional_formatting.add(rng, ColorScaleRule(
            start_type="min", start_color="F8696B",
            mid_type="percentile", mid_value=50, mid_color="FFEB84",
            end_type="max", end_color="63BE7B",
        ))
    if "recommendation" in cols:
        letter = get_column_letter(cols.index("recommendation") + 1)
        rng = f"{letter}2:{letter}{n_rows + 1}"
        ws.conditional_formatting.add(rng, CellIsRule(
            operator="equal", formula=['"select"'], fill=CF_SELECT_FILL, font=CF_SELECT_FONT))
        ws.conditional_formatting.add(rng, CellIsRule(
            operator="equal", formula=['"monitor"'], fill=CF_MONITOR_FILL, font=CF_MONITOR_FONT))


def recalibrate_recommendations(df: pd.DataFrame) -> pd.DataFrame:
    """بازکالیبره کردن توصیه‌ها بر اساس صدک نمره."""
    print("Recalibrating recommendations with percentile method...")

    scores = df["export_score"]
    select_threshold = float(np.percentile(scores, SELECT_PERCENTILE))
    monitor_threshold = float(np.percentile(scores, MONITOR_PERCENTILE))

    print(f"  Score range: [{scores.min():.3f}, {scores.max():.3f}]")
    print(f"  Select threshold (percentile {SELECT_PERCENTILE}): {select_threshold:.3f}")
    print(f"  Monitor threshold (percentile {MONITOR_PERCENTILE}): {monitor_threshold:.3f}")

    def new_rec(row):
        score = row["export_score"]
        has_growth = (row.get("n_strong_growth", 0) > 0 or
                      row.get("n_moderate_growth", 0) > 0 or
                      row.get("n_emerging", 0) > 0)
        if score >= select_threshold and has_growth:
            return "select"
        elif score >= monitor_threshold:
            return "monitor"
        else:
            return "investigate"

    df["recommendation_old"] = df["recommendation"]
    df["recommendation"] = df.apply(new_rec, axis=1)
    df["select_threshold"] = select_threshold
    df["monitor_threshold"] = monitor_threshold

    # آمار
    print("\n  Distribution after recalibration:")
    for rec in ("select", "monitor", "investigate"):
        count = int((df["recommendation"] == rec).sum())
        print(f"    {rec}: {count} ({count / len(df) * 100:.1f}%)")
    assert set(df["recommendation"].unique()) <= {"select", "monitor", "investigate"}
    return df


def build_country_summary(df: pd.DataFrame, pair_values: pd.DataFrame) -> pd.DataFrame:
    """Top 30 کشور مقصد بر اساس تعداد HS Codeهای رشد (حضور در top_5 کشورهای هر کاندید).

    ارزش ۱۴۰۴ هر کشور = جمع ارزش واقعی جفت‌های HS×Country از trend-analysis-5y.parquet.
    """
    records = []
    for _, row in df.iterrows():
        raw = row.get("top_5_countries")
        if raw is None or pd.isna(raw):
            continue
        for entry in str(raw).split(";"):
            iso2 = entry.split("(")[0].strip()
            if iso2:
                records.append({
                    "hs_code": row["hs_code"],
                    "iso2": iso2,
                    "recommendation": row["recommendation"],
                })
    pairs = pd.DataFrame(records)

    merged = pairs.merge(
        pair_values,
        left_on=["hs_code", "iso2"],
        right_on=["hs_code", "destination_country_iso2"],
        how="left",
    )
    n_missing = int(merged["value_1404"].isna().sum())
    if n_missing:
        print(f"  ⚠️ {n_missing} جفت در trend parquet یافت نشد — ارزش صفر در نظر گرفته شد")
        merged["value_1404"] = merged["value_1404"].fillna(0.0)

    summary = (
        merged.groupby(["iso2", "destination_country_fa"], dropna=False)
        .agg(
            n_hs=("hs_code", "count"),
            n_select=("recommendation", lambda x: int((x == "select").sum())),
            n_monitor=("recommendation", lambda x: int((x == "monitor").sum())),
            total_value_1404=("value_1404", "sum"),
        )
        .reset_index()
        .rename(columns={"iso2": "destination_country_iso2"})
    )
    summary = summary[
        ["destination_country_iso2", "destination_country_fa",
         "n_hs", "n_select", "n_monitor", "total_value_1404"]
    ]
    return (
        summary
        .sort_values(["n_hs", "total_value_1404"], ascending=[False, False])
        .head(30)
        .reset_index(drop=True)
    )


def build_excel(df: pd.DataFrame, pair_values: pd.DataFrame) -> None:
    """ساخت فایل Excel کامل با ۷ شیت."""
    print("\nBuilding Excel file...")

    wb = Workbook()
    wb.remove(wb.active)
    wb.properties.creator = "Z.ai"
    wb.properties.title = "Iran Export Candidates 1400-1404 (task-12 — recalibrated)"

    # ─── شیت ۱: All-Candidates ───
    print("  Sheet 1/7: All-Candidates")
    ws1 = wb.create_sheet("All-Candidates")
    n1 = write_table(ws1, df, EXCEL_COLS)
    add_conditional_formatting(ws1, EXCEL_COLS, n1)
    print(f"    {n1} rows")

    # ─── شیت ۲: Top-500 ───
    print("  Sheet 2/7: Top-500")
    ws2 = wb.create_sheet("Top-500")
    n2 = write_table(ws2, df.head(500), EXCEL_COLS)
    add_conditional_formatting(ws2, EXCEL_COLS, n2)
    print(f"    {n2} rows")

    # ─── شیت ۳: Select ───
    print("  Sheet 3/7: Select")
    ws3 = wb.create_sheet("Select")
    sel_df = df[df["recommendation"] == "select"]
    n3 = write_table(ws3, sel_df, EXCEL_COLS)
    add_conditional_formatting(ws3, EXCEL_COLS, n3)
    print(f"    {n3} rows")

    # ─── شیت ۴: Monitor ───
    print("  Sheet 4/7: Monitor")
    ws4 = wb.create_sheet("Monitor")
    mon_df = df[df["recommendation"] == "monitor"]
    n4 = write_table(ws4, mon_df, EXCEL_COLS)
    add_conditional_formatting(ws4, EXCEL_COLS, n4)
    print(f"    {n4} rows")

    # ─── شیت ۵: By-Chapter ───
    print("  Sheet 5/7: By-Chapter")
    chapter_stats = (
        df.groupby("hs_chapter")
        .agg(
            n_candidates=("hs_code", "count"),
            n_select=("recommendation", lambda x: int((x == "select").sum())),
            n_monitor=("recommendation", lambda x: int((x == "monitor").sum())),
            total_score=("export_score", "sum"),
            avg_score=("export_score", "mean"),
            max_score=("export_score", "max"),
            total_value_1404=("total_value_1404", "sum"),
        )
        .reset_index()
        .sort_values("total_score", ascending=False)
        .reset_index(drop=True)
    )
    ws5 = wb.create_sheet("By-Chapter")
    n5 = write_table(ws5, chapter_stats, list(chapter_stats.columns))
    print(f"    {n5} chapters")

    # ─── شیت ۶: By-Target-Country ───
    print("  Sheet 6/7: By-Target-Country")
    country_summary = build_country_summary(df, pair_values)
    ws6 = wb.create_sheet("By-Target-Country")
    n6 = write_table(ws6, country_summary, list(country_summary.columns))
    print(f"    {n6} countries (Top 30 by n_hs)")

    # ─── شیت ۷: Methodology ───
    print("  Sheet 7/7: Methodology")
    build_methodology_sheet(wb, df)

    OUTPUT_EXCEL.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUTPUT_EXCEL)
    print(f"  Saved: {OUTPUT_EXCEL}")
    print(f"  File size: {OUTPUT_EXCEL.stat().st_size / 1024:.1f} KB")


def build_methodology_sheet(wb: Workbook, df: pd.DataFrame) -> None:
    """شیت ۷: توضیح متدولوژی + وزن‌ها + آستانه‌های بازکالیبره‌شده."""
    ws = wb.create_sheet("Methodology")

    select_thresh = float(df["select_threshold"].iloc[0])
    monitor_thresh = float(df["monitor_threshold"].iloc[0])
    n_sel = int((df["recommendation"] == "select").sum())
    n_mon = int((df["recommendation"] == "monitor").sum())
    n_inv = int((df["recommendation"] == "investigate").sum())

    methodology = [
        ["گزارش کاندیدهای صادرات ایران (1400-1404)", ""],
        ["", ""],
        ["تاریخ تولید", TODAY],
        ["منبع داده", "service.tccim.ir (آینه گمرک)"],
        ["سال‌های موجود", "1400, 1401, 1402, 1403, 1404"],
        ["سال 1405", "موجود نیست (Issue-008)"],
        ["", ""],
        ["─── متدولوژی نمره‌دهی ───", ""],
        ["export_score = 0.30×norm(cagr_5y) + 0.20×norm(slope) + 0.20×norm(mean_recent_3y)", ""],
        ["             + 0.10×norm(n_destinations_active) + 0.10×norm(trend_consistency)", ""],
        ["             - 0.10×norm(cv)", ""],
        ["", ""],
        ["─── وزن‌ها (Decision-009) ───", ""],
        ["w1 (CAGR)", 0.30],
        ["w2 (slope)", 0.20],
        ["w3 (حجم)", 0.20],
        ["w4 (تنوع بازار)", 0.10],
        ["w5 (ثبات)", 0.10],
        ["w6 (جریمه نوسان)", 0.10],
        ["", ""],
        ["─── آستانه‌های بازکالیبره‌شده (صدک) ───", ""],
        ["روش", "Percentile-based (نسبی)"],
        ["چرا بازکالیبره شد؟", "دامنه واقعی نمره‌ها [-0.03, 0.35] بود؛ آستانه‌های ثابت 0.4/0.7 همه را investigate می‌کردند"],
        ["select threshold (percentile 90)", select_thresh],
        ["monitor threshold (percentile 60)", monitor_thresh],
        ["توجه", "آستانه‌های صدک نسبی‌اند — با تغییر داده/فیلترها باید دوباره محاسبه شوند"],
        ["", ""],
        ["─── شیت‌ها ───", ""],
        ["All-Candidates", f"همه {len(df)} کاندید"],
        ["Top-500", "۵۰۰ کاندید اول بر اساس رتبه"],
        ["Select", f"{n_sel} کاندید select (نمره ≥ صدک ۹۰)"],
        ["Monitor", f"{n_mon} کاندید monitor (صدک ۶۰ تا ۹۰)"],
        ["By-Chapter", "تجمیع به ازای فصل HS"],
        ["By-Target-Country", "Top 30 کشور مقصد — ارزش واقعی سطح-جفت از trend-analysis-5y.parquet (task-10)"],
        ["Methodology", "این شیت"],
        ["", ""],
        ["─── فیلتر اولیه (روش membership) ───", ""],
        ["شرط 1", "n_strong_growth > 0 OR n_moderate_growth > 0 OR n_emerging > 0"],
        ["شرط 2", "mean_recent_3y_aggregated > 1,000,000 USD"],
        ["شرط 3", "n_destinations_active >= 3"],
        ["", ""],
        ["─── آمار کلی ───", ""],
        ["تعداد کل کاندیدها", len(df)],
        ["تعداد select", n_sel],
        ["تعداد monitor", n_mon],
        ["تعداد investigate", n_inv],
        ["دامنه نمره", f"[{df['export_score'].min():.3f}, {df['export_score'].max():.3f}]"],
        ["", ""],
        ["─── ریسک‌ها ───", ""],
        ["volatile", "cv_aggregated > 0.7 (با صفر-پرکردن جفت‌ها متورم است — همه ۵۷۴ فعال)"],
        ["concentrated", "سهم بزرگ‌ترین کشور > 70٪"],
        ["declining_recent", "value_1404 < value_1403"],
        ["", ""],
        ["─── محدودیت‌ها ───", ""],
        ["1", "سال 1405 موجود نیست (Issue-008)"],
        ["2", "تفکیک ماهانه 1403/1404 موجود نیست (Issue-006)"],
        ["3", "Mann-Kendall با n=5: حداقل p ≈ 0.0275"],
        ["4", "آستانه‌های صدک نسبی هستند، نه مطلق"],
    ]

    for row_idx, (key, val) in enumerate(methodology, 1):
        is_section = str(key).startswith("───")
        kcell = ws.cell(row=row_idx, column=1, value=key)
        kcell.font = Font(
            name="Calibri", size=11, bold=not is_section,
            color="1F4E78" if is_section else "37352F",
        )
        vcell = ws.cell(row=row_idx, column=2, value=_native_value(val))
        if isinstance(val, float):
            vcell.number_format = "0.000"
        elif isinstance(val, int):
            vcell.number_format = "#,##0"

    ws.column_dimensions["A"].width = 60
    ws.column_dimensions["B"].width = 55


def update_executive_ranking(df: pd.DataFrame) -> None:
    """به‌روزرسانی گزارش اجرایی با Top 50 + آمار بازکالیبره."""
    print("\nUpdating executive ranking with Top 50...")

    top_50 = df.head(50)
    select_count = int((df["recommendation"] == "select").sum())
    monitor_count = int((df["recommendation"] == "monitor").sum())
    investigate_count = int((df["recommendation"] == "investigate").sum())

    select_thresh = float(df["select_threshold"].iloc[0])
    monitor_thresh = float(df["monitor_threshold"].iloc[0])

    content = f"""---
type: analysis
title: 🎯 گزارش اجرایی کاندیدهای صادرات (task-12 — بازکالیبره‌شده)
created: 2026-09-14
updated: {TODAY}
status: review
tags:
  - analysis
  - export-candidates
  - ranking
  - recalibrated
folder: 06-Analysis/export-candidates
---

# 🎯 گزارش اجرایی کاندیدهای صادرات (task-12 — بازکالیبره‌شده)

**تاریخ به‌روزرسانی**: {TODAY} ({TODAY_FA})
**Agent**: Analyst
**Pipeline**: `scripts/recalibrate_and_export_excel.py`
**منبع داده**: `export-candidates-ranked.parquet` (task-12) + `trend-analysis-5y.parquet` (task-10، برای ارزش‌های سطح-جفت)

## ۱. خلاصه اجرایی

این گزارش نسخه بازکالیبره‌شده task-12 است. نسخه قبلی آستانه‌های ثابت ۰.۴/۰.۷ داشت که با دامنه واقعی نمره‌ها [-0.03, 0.35] همخوانی نداشت — همه ۵۷۴ کاندید «investigate» شده بودند. در این نسخه، آستانه‌ها بر اساس **صدک** محاسبه شده‌اند تا توصیه‌ها نسبی و قابل‌استفاده باشند.

### متدولوژی نمره‌دهی (بدون تغییر)

```
export_score = 0.30 × norm(cagr_5y)
             + 0.20 × norm(slope)
             + 0.20 × norm(mean_recent_3y)
             + 0.10 × norm(n_destinations_active)
             + 0.10 × norm(trend_consistency)
             - 0.10 × norm(cv)
```

⚠️ **تطابق با Issue-008**: از cagr_5y به‌جای cagr_6y استفاده شده (سال 1405 موجود نیست).

### آستانه‌های بازکالیبره‌شده (صدک)

| توصیه | آستانه | تعداد | درصد |
|-------|--------|-------|-------|
| **select** | نمره ≥ {select_thresh:.3f} (صدک ۹۰) | {select_count} | {select_count / len(df) * 100:.1f}% |
| **monitor** | نمره ≥ {monitor_thresh:.3f} و < {select_thresh:.3f} (صدک ۶۰-۹۰) | {monitor_count} | {monitor_count / len(df) * 100:.1f}% |
| **investigate** | نمره < {monitor_thresh:.3f} (صدک ۶۰) | {investigate_count} | {investigate_count / len(df) * 100:.1f}% |

### فیلتر اولیه (روش membership)

طبق یافته task-11، فیلتر dominant_trend کار نمی‌کرد. روش membership استفاده شد:
- HS Codeهایی که حداقل یک جفت در strong_growth، moderate_growth، یا emerging دارند.
- mean_recent_3y_aggregated > 1,000,000 USD.
- n_destinations_active >= 3.

**نتیجه فیلتر**: {len(df):,} HS Code کاندید.

## ۲. Top 50 کاندید صادرات

| رتبه | HS Code | شرح کالا | نمره | توصیه | CAGR 5y | حجم ۳ سال اخیر (USD) | کشورهای هدف |
|------|---------|----------|------|-------|---------|----------------------|-------------|
"""

    for _, row in top_50.iterrows():
        cagr = row.get("cagr_5y_aggregated")
        cagr_str = f"{cagr * 100:.1f}%" if pd.notna(cagr) else "N/A"
        vol = row.get("mean_recent_3y_aggregated")
        vol_str = f"{vol / 1e6:,.0f}M" if pd.notna(vol) and vol > 0 else "N/A"
        desc = str(row["hs_description"])[:35]
        countries = str(row.get("top_5_countries", ""))[:50]

        content += (
            f"| {row['rank']} | {row['hs_code']} | {desc} | {row['export_score']:.3f} "
            f"| {row['recommendation']} | {cagr_str} | {vol_str} | {countries} |\n"
        )

    # تحلیل کیفی Top 5
    content += "\n## ۳. تحلیل کیفی Top 5\n\n"

    for _, row in top_50.head(5).iterrows():
        cagr = row.get("cagr_5y_aggregated")
        cagr_str = f"{cagr * 100:.1f}%" if pd.notna(cagr) else "N/A"
        vol = row.get("mean_recent_3y_aggregated")
        vol_str = f"{vol / 1e6:,.0f}" if pd.notna(vol) else "N/A"
        rf = row.get("risk_flags")
        rf_str = str(rf) if pd.notna(rf) else "none"

        content += f"""### رتبه {row['rank']}: {row['hs_description']} ({row['hs_code']})

- **نمره**: {row['export_score']:.3f}
- **توصیه**: {row['recommendation']}
- **CAGR ۵ ساله**: {cagr_str}
- **حجم ۳ سال اخیر**: {vol_str} میلیون دلار
- **تعداد کشورهای فعال**: {row.get('n_destinations_active', 0)}
- **ریسک‌ها**: {rf_str}
- **کشورهای هدف**: {row.get('top_5_countries', '')}

"""

    # آمار فصل‌ها
    content += """## ۴. فصل‌های پرپتانسیل (Top 10)

| فصل | تعداد کاندید | select | monitor | مجموع نمره | میانگین نمره |
|------|---------------|--------|---------|------------|--------------|
"""

    chapter_stats = (
        df.groupby("hs_chapter")
        .agg(
            n_candidates=("hs_code", "count"),
            n_select=("recommendation", lambda x: int((x == "select").sum())),
            n_monitor=("recommendation", lambda x: int((x == "monitor").sum())),
            total_score=("export_score", "sum"),
            avg_score=("export_score", "mean"),
        )
        .sort_values("total_score", ascending=False)
        .head(10)
    )

    for chapter, stats in chapter_stats.iterrows():
        # iterrows سطر را به float تبدیل می‌کند — int() برای نمایش صحیح شمارنده‌ها
        content += (
            f"| {chapter} | {int(stats['n_candidates'])} | {int(stats['n_select'])} | {int(stats['n_monitor'])} "
            f"| {stats['total_score']:.2f} | {stats['avg_score']:.3f} |\n"
        )

    content += f"""

## ۵. یادآوری: چرا آستانه‌های ثابت کار نکردند

در نسخه اول task-12، آستانه‌های ثابت ۰.۴/۰.۷ تعریف شده بود. اما دامنه واقعی نمره‌ها [-0.03, 0.35] بود — هیچ محصولی هم‌زمان در هر ۶ بُعد (CAGR + slope + حجم + تنوع + ثبات - نوسان) پیشتاز نبود. این طبیعی است چون:

1. **محدودیت داده**: با ۵ سال داده و محدودیت‌های منبع، شاخص‌ها هم‌زمان نمی‌توانند همه عالی باشند.
2. **جریمه نوسان (cv)**: هرچه محصول پرحجم‌تر، نوسان کمتر اما رشد ممکن است کمتر باشد.
3. **تنوع در برابر تمرکز**: محصولات با تنوع بالا معمولاً حجم کمتر دارند.

**راه‌حل**: آستانه‌های نسبی (صدک) — توصیه نسبت به سایر کاندیدها، نه مطلق.

## ۶. فایل‌های خروجی

- **Excel کامل**: `07-Exports/iran-export-candidates-500.xlsx` (۷ شیت)
  - All-Candidates (همه {len(df)} کاندید؛ conditional formatting روی export_score + رنگ وضعیت توصیه)
  - Top-500 (۵۰۰ کاندید اول بر اساس رتبه)
  - Select ({select_count} کاندید — صدک ۹۰)
  - Monitor ({monitor_count} کاندید — صدک ۶۰-۹۰)
  - By-Chapter (تجمیع {df['hs_chapter'].nunique()} فصل HS)
  - By-Target-Country (Top 30 کشور مقصد — ارزش واقعی سطح-جفت HS×Country از `trend-analysis-5y.parquet`، نه انتساب تکراری)
  - Methodology (متدولوژی + وزن‌ها + آستانه‌ها)
- **Parquet بازکالیبره‌شده**: `05-Data/processed/export-candidates-ranked-recalibrated.parquet` (ستون‌های جدید: `recommendation_old`، `select_threshold`، `monitor_threshold`)
- **این گزارش**: `06-Analysis/export-candidates/_executive-ranking.md`

## ۷. محدودیت‌ها

1. سال 1405: داده موجود نیست (Issue-008).
2. تفکیک ماهانه 1403/1404: موجود نیست (Issue-006).
3. هشدار آماری MK: با n=5، حداقل p-value ≈ 0.0275.
4. آستانه‌های صدک نسبی هستند — با تغییر داده، آستانه‌ها هم تغییر می‌کنند.

---

**این گزارش هدف نهایی پروژه است.** برای تغییر وزن‌ها یا آستانه‌ها، با کاربر هماهنگ کنید.
"""

    with open(OUTPUT_EXEC, "w", encoding="utf-8") as f:
        f.write(content)
    print(f"  Saved: {OUTPUT_EXEC}")


def main():
    print("=" * 60)
    print("task-12 (افزونه): بازکالیبره آستانه‌ها (صدکی) + Excel کامل")
    print("=" * 60)

    print("\n[1/5] Loading data...")
    df = pd.read_parquet(INPUT_PATH)
    print(f"  Loaded: {len(df)} candidates × {df.shape[1]} columns")

    print("\n[2/5] Loading pair-level values (for By-Target-Country sheet)...")
    pair_values = pd.read_parquet(
        TREND_PATH,
        columns=["hs_code", "destination_country_iso2", "destination_country_fa", "value_1404"],
    )
    print(f"  Loaded: {len(pair_values):,} HS×Country pairs from trend-analysis-5y.parquet")

    print("\n[3/5] Recalibrating recommendations...")
    df = recalibrate_recommendations(df)
    df = df.sort_values("rank").reset_index(drop=True)

    print("\n[4/5] Saving recalibrated Parquet...")
    pq.write_table(pa.Table.from_pandas(df, preserve_index=False), OUTPUT_PARQUET, compression="snappy")
    print(f"  Saved: {OUTPUT_PARQUET} ({OUTPUT_PARQUET.stat().st_size / 1024:.1f} KB)")

    print("\n[5/5] Building Excel + updating executive ranking...")
    build_excel(df, pair_values)
    update_executive_ranking(df)

    print("\n" + "=" * 60)
    print("Top 10 candidates (recalibrated):")
    for _, row in df.head(10).iterrows():
        cagr = row.get("cagr_5y_aggregated")
        cagr_str = f"CAGR {cagr * 100:.0f}%" if pd.notna(cagr) else "CAGR N/A"
        print(
            f"  {row['rank']:>3}. {row['hs_code']} — {str(row['hs_description'])[:45]} "
            f"(score={row['export_score']:.3f}, {row['recommendation']}, {cagr_str})"
        )
    print("=" * 60)
    print(f"\n✅ Done!")
    print(f"   Excel: {OUTPUT_EXCEL}")
    print(f"   Executive ranking updated: {OUTPUT_EXEC}")
    print(f"   Recalibrated parquet: {OUTPUT_PARQUET}")


if __name__ == "__main__":
    main()
